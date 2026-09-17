import streamlit as st
import sqlite3
import pandas as pd
import io
import base64
from datetime import datetime, timedelta
from core.database import get_connection
from core.documentos import formatar_data_br, normalizar_data, exportar_excel_seguro, gerar_pdf_tabela_seguro
from core.api_cnpj import buscar_dados_cnpj
from core.gestor_fluxo import GestorFluxoMilitar
from core.documentos import gerar_requisicao_odt, gerar_requisicao_docx
try:
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    docx = None


def render(user):
    st.title("📝 Elaboração de Processos de Aquisição")
    st.info("Selecione o tipo de processo e o modelo para carregar os campos específicos e gerar o documento com formatação oficial.")
    conn = get_connection()

    aba_proc1, aba_proc2 = st.tabs(["📄 Elaborar Processo por Modelo", "📁 Banco de Modelos (LibreOffice / Word)"])

    with aba_proc1:
        st.markdown("### 1. Seleção do Processo e Modelo")
        col_m1, col_m2 = st.columns(2)

        tipo_proc_sel = col_m1.selectbox(
            "Modalidade / Tipo de Processo:",
            [
                "Pregão Eletrônico (Próprio)",
                "Pregão Eletrônico (Participante)",
                "Adesão à Ata de Registro de Preços (Carona)",
                "Dispensa de Licitação (Art. 75 Lei 14.133/21)",
                "Inexigibilidade de Licitação"
            ],
            key="sel_tipo_proc_box"
        )

        c = conn.cursor()
        c.execute('''
        SELECT id, nome_modelo, tipo_processo, formato, arquivo_docx 
        FROM processos_modelos 
        WHERE om_id = ? AND (tipo_processo = ? OR tipo_processo LIKE '%Requisição%' OR tipo_processo LIKE '%Pregão%')
        ''', (user['om_id'], tipo_proc_sel))
        modelos_banco = c.fetchall()

        opcoes_modelos_dict = {"02. Requisição de Despesa (Modelo Oficial - B Adm Ap / 5ª RM)": None}
        for m_row in modelos_banco:
            opcoes_modelos_dict[f"{m_row['nome_modelo']} ({m_row['formato'].upper()})"] = m_row['id']

        modelo_escolhido_nome = col_m2.selectbox(
            "Documento / Modelo a Elaborar:",
            list(opcoes_modelos_dict.keys()),
            key="sel_modelo_doc_box"
        )
        modelo_id_selecionado = opcoes_modelos_dict[modelo_escolhido_nome]

        st.markdown("---")

        st.markdown("### 2. Dados da Requisição")
        st.caption("🔒 O cabeçalho oficial da Base de Administração e Apoio da 5ª RM é fixo e inalterável.")

        col_r1, col_r2, col_r3 = st.columns(3)
        num_req = col_r1.text_input("Número da Requisição:", value="5-Seção de Contratações/B Adm Ap")
        nup_req = col_r2.text_input("NUP / Número do EB:", value="65378.006979/2026-34")
        assunto_req = col_r3.text_input("Assunto da Requisição:", value="aquisição de material de expediente (PRÓPRIO)")

        data_extenso_req = st.text_input("Local e Data por Extenso:", value="Curitiba, PR, 22 de junho de 2026")

        st.markdown("### 3. Demonstrativo de Necessidades & Fornecedor")
        col_p1, col_p2 = st.columns(2)
        pregao_req = col_p1.text_input("Identificação do Pregão / Processo de Origem:", value="Pregão Eletrônico 90001/2026, da B Adm Ap /5ª RM (UG 160192)")
        tipo_emp_req = col_p2.selectbox("Tipo de Empenho:", ["Ordinário", "Estimativo", "Global"])

        col_cnpj_p1, col_cnpj_p2 = st.columns([3, 1])
        cnpj_forn_req = col_cnpj_p1.text_input("CNPJ do Fornecedor:", value="09.247.343/0001-20")
        if col_cnpj_p2.button("🔍 Buscar Receita", use_container_width=True):
            d_emp = buscar_dados_cnpj(cnpj_forn_req, conn)
            if d_emp:
                st.session_state['req_forn_nome'] = d_emp.get('razao_social', '')
                st.success(f"Empresa: {st.session_state['req_forn_nome']}")

        razao_social_req = st.text_input("Razão Social do Fornecedor:", value=st.session_state.get('req_forn_nome', "BIOLIMP SERVICOS ESPECIALIZADOS DE HIGIENIZACAO TEXTIL EIRELI"))

        st.markdown("### 4. Tabela de Itens da Requisição")
        if 'itens_req_df' not in st.session_state:
            st.session_state['itens_req_df'] = pd.DataFrame([
                {"Item": 1, "Descrição": "COPO DESCARTAVEL 200 ML", "UN": "CX", "QTD": 2.0, "V_Unit": 92.00},
                {"Item": 2, "Descrição": "DETERGENTE LIQUIDO NEUTRO", "UN": "UN", "QTD": 78.0, "V_Unit": 1.18}
            ])

        df_edit_req = st.data_editor(
            st.session_state['itens_req_df'],
            num_rows="dynamic",
            use_container_width=True,
            column_config={
                "Item": st.column_config.NumberColumn("Item Nº", min_value=1, step=1),
                "Descrição": st.column_config.TextColumn("Descrição Detalhada", width="large", required=True),
                "UN": st.column_config.TextColumn("UN", width="small", default="UN"),
                "QTD": st.column_config.NumberColumn("QTD", min_value=1.0, step=1.0),
                "V_Unit": st.column_config.NumberColumn("Valor Unitário (R$)", min_value=0.01, step=0.1, format="R$ %.2f")
            }
        )
        st.session_state['itens_req_df'] = df_edit_req
        st.caption("ℹ️ * Inserir os dados conforme consta no Termo de Homologação.")

        df_calc_req = df_edit_req.copy()
        df_calc_req['Total'] = df_calc_req['QTD'] * df_calc_req['V_Unit']
        total_requisicao = df_calc_req['Total'].sum()

        st.markdown(f'''
        <div style="background: rgba(34, 197, 94, 0.15); border: 1.5px solid #22c55e; border-radius: 8px; padding: 12px 18px; display: flex; justify-content: space-between; align-items: center; margin: 15px 0;">
            <span style="font-size: 15px; font-weight: bold; color: #fff;">VALOR TOTAL DA REQUISIÇÃO:</span>
            <span style="font-size: 22px; font-weight: 900; color: #22c55e;">R$ {total_requisicao:,.2f}</span>
        </div>
        ''', unsafe_allow_html=True)

        st.markdown("### 5. Local de Entrega e Justificativas da Despesa")
        local_entrega_req = st.text_area(
            "3. Local de Entrega dos Materiais / Prestação dos Serviços:",
            value="Base de Administração e Apoio da 5ª Região Militar, situado à Rua 31 de Março, S/Nr, CEP 81.150-900, Bairro Pinheirinho, em Curitiba, PR.",
            height=70
        )

        col_j1, col_j2 = st.columns(2)
        just_a = col_j1.text_area(
            "4. a. Necessidade (Problema a ser resolvido):",
            value="a contratação/aquisição é necessária para garantir o funcionamento administrativo e operacional desta OM, visto que os estoques atuais se esgotaram.",
            height=100
        )
        just_b = col_j2.text_area(
            "4. b. Consequência da não emissão (Prejuízo potencial):",
            value="a não emissão desta NE acarretará prejuízo ao funcionamento regular da OM e à execução de suas atividades administrativas e operacionais.",
            height=100
        )

        st.markdown("### 6. Classificação Orçamentária")
        df_ncs_disp = pd.read_sql_query(f"SELECT id, numero_nc, pi, natureza_despesa FROM notas_credito WHERE om_id = {user['om_id']}", conn)
        nc_selec_req = None
        pi_padrao = "IXAPFUNADOM"
        nd_padrao = "339030"

        if not df_ncs_disp.empty:
            dict_nc_opcoes = {f"{r['numero_nc']} (PI: {r['pi']})": r for _, r in df_ncs_disp.iterrows()}
            escolha_nc = st.selectbox("Puxar Dados da Nota de Crédito da OM:", list(dict_nc_opcoes.keys()))
            nc_escolhida_row = dict_nc_opcoes[escolha_nc]
            nc_selec_req = nc_escolhida_row['numero_nc']
            pi_padrao = nc_escolhida_row['pi'] or pi_padrao
            nd_padrao = nc_escolhida_row['natureza_despesa'] or nd_padrao

        col_orc1, col_orc2, col_orc3, col_orc4 = st.columns(4)
        pi_req = col_orc1.text_input("Plano Interno (PI):", value=pi_padrao)
        ptres_req = col_orc2.text_input("PTRES:", value="171502")
        nd_req = col_orc3.text_input("Natureza da Despesa (ND):", value=nd_padrao)
        nc_num_req = col_orc4.text_input("Número da NC:", value=nc_selec_req if nc_selec_req else "2026NC000672")

        st.markdown("### 7. Assinatura do Fiscal Administrativo")
        col_fisc1, col_fisc2 = st.columns(2)
        fisc_nome = col_fisc1.text_input("Nome Completo e Posto/Graduação:", value="MARCO AURÉLIO GOBETTI DA FONSECA - Cap")
        fisc_funcao = col_fisc2.text_input("Função do Fiscal:", value="Fiscal Administrativo - B Adm Ap/5ª RM")

        st.markdown("---")

        st.markdown("### 👁️ Pré-Visualização da Folha Oficial (Padrão A4)")

        linhas_tabela_html = ""
        for it in df_calc_req.to_dict('records'):
            linhas_tabela_html += f"""
            <tr style="border-bottom: 1px solid #111;">
                <td style="padding: 6px; text-align: center; border-right: 1px solid #111;">{it['Item']}</td>
                <td style="padding: 6px; border-right: 1px solid #111;">{it['Descrição']}</td>
                <td style="padding: 6px; text-align: center; border-right: 1px solid #111;">{it['UN']}</td>
                <td style="padding: 6px; text-align: center; border-right: 1px solid #111;">{it['QTD']:.0f}</td>
                <td style="padding: 6px; text-align: right; border-right: 1px solid #111;">R$ {it['V_Unit']:.2f}</td>
                <td style="padding: 6px; text-align: right;">R$ {it['Total']:.2f}</td>
            </tr>
            """

        st.markdown(f"""
        <div style="background: #ffffff; color: #111827; padding: 45px 50px; border-radius: 4px; box-shadow: 0 10px 25px rgba(0,0,0,0.5); max-width: 820px; margin: 15px auto; font-family: 'Times New Roman', Times, serif; font-size: 14px; line-height: 1.4; border: 1px solid #d1d5db;">
            <div style="text-align: center; margin-bottom: 15px;">
                <img src="https://upload.wikimedia.org/wikipedia/commons/thumb/b/bf/Coat_of_arms_of_Brazil.svg/150px-Coat_of_arms_of_Brazil.svg.png" style="height: 72px; width: auto; margin-bottom: 6px;">
                <div style="font-weight: bold; font-size: 13px; line-height: 1.25; color: #000;">
                    MINISTÉRIO DA DEFESA<br>
                    EXÉRCITO BRASILEIRO<br>
                    BASE DE ADMINISTRAÇÃO E APOIO DA 5ª REGIÃO MILITAR<br>
                    <span style="font-weight: normal; font-size: 12px;">(Companhia do QG da 5ª RM/DI)</span><br>
                    <span style="font-weight: bold; font-size: 12px;">BASE MAJOR AGOSTINHO JOSÉ RODRIGUES</span>
                </div>
            </div>

            <div style="margin-top: 20px;">
                <b>Requisição Nº</b> {num_req}<br>
                <b>EB:</b> {nup_req}<br>
                <b>Assunto:</b> {assunto_req}<br>
                <b>Anexos:</b> Termo de Homologação, Nota de Crédito, CADIN, SICAF, TCU
            </div>

            <div style="text-align: right; margin: 15px 0;">
                Curitiba, PR, {data_extenso_req}.
            </div>

            <p style="text-align: justify; text-indent: 30px;">
                1. Nos termos contidos no Art. 13 das IG 12-02, aprovadas pela Portaria Ministerial nº 305, de 24 MAIO 95, solicito providências no sentido de aprovar a seguinte despesa com aquisição do material abaixo especificado, a fim de atender necessidades desta OM.
            </p>

            <p><b>2. DEMONSTRATIVO DE NECESSIDADES:</b><br>
            {pregao_req}<br>
            <b>RAZÃO SOCIAL:</b> {razao_social_req} — <b>CNPJ:</b> {cnpj_forn_req}</p>

            <table style="width: 100%; border-collapse: collapse; border: 1.5px solid #111; font-size: 12px; margin: 10px 0;">
                <thead>
                    <tr style="background-color: #f3f4f6; border-bottom: 1.5px solid #111;">
                        <th style="padding: 6px; border-right: 1px solid #111; width: 8%;">Item</th>
                        <th style="padding: 6px; border-right: 1px solid #111; text-align: left;">Descrição</th>
                        <th style="padding: 6px; border-right: 1px solid #111; width: 8%;">UN</th>
                        <th style="padding: 6px; border-right: 1px solid #111; width: 8%;">QTD</th>
                        <th style="padding: 6px; border-right: 1px solid #111; width: 18%; text-align: right;">VALOR UNITÁRIO</th>
                        <th style="padding: 6px; width: 18%; text-align: right;">TOTAL</th>
                    </tr>
                </thead>
                <tbody>
                    {linhas_tabela_html}
                </tbody>
            </table>

            <div style="text-align: right; font-weight: bold; font-size: 13px; margin: 5px 0 10px 0;">
                VALOR TOTAL: R$ {total_requisicao:,.2f}
            </div>

            <p><b>TIPO DE EMPENHO:</b> ({'X' if tipo_emp_req=='Ordinário' else ' '}) Ordinário &nbsp;&nbsp;&nbsp;&nbsp; ({'X' if tipo_emp_req=='Estimativo' else ' '}) Estimativo &nbsp;&nbsp;&nbsp;&nbsp; ({'X' if tipo_emp_req=='Global' else ' '}) Global</p>

            <p><b>3. ENTREGA DO(S) BEM(NS)/PRESTAÇÃO DO(S) SERVIÇO(S):</b><br>
            Os materiais serão entregues no seguinte local: {local_entrega_req}</p>

            <p><b>4. JUSTIFICATIVA DA DESPESA:</b><br>
            <b>a. Necessidade:</b> {just_a}<br>
            <b>b. Consequência da não emissão:</b> {just_b}</p>

            <p><b>5. CLASSIFICAÇÃO ORÇAMENTÁRIA:</b></p>
            <table style="width: 100%; border-collapse: collapse; border: 1px solid #111; font-size: 12px; margin-bottom: 15px;">
                <tr style="background: #f3f4f6; border-bottom: 1px solid #111; text-align: center; font-weight: bold;">
                    <td style="padding: 5px; border-right: 1px solid #111;">PI</td>
                    <td style="padding: 5px; border-right: 1px solid #111;">PTRES</td>
                    <td style="padding: 5px; border-right: 1px solid #111;">ND</td>
                    <td style="padding: 5px;">NC</td>
                </tr>
                <tr style="text-align: center;">
                    <td style="padding: 5px; border-right: 1px solid #111;">{pi_req}</td>
                    <td style="padding: 5px; border-right: 1px solid #111;">{ptres_req}</td>
                    <td style="padding: 5px; border-right: 1px solid #111;">{nd_req}</td>
                    <td style="padding: 5px;">{nc_num_req}</td>
                </tr>
            </table>

            <div style="text-align: center; margin-top: 35px;">
                Aquisição aprovada por:<br><br>
                Curitiba, PR, {data_extenso_req}.<br><br><br>
                ____________________________________________<br>
                <b>{fisc_nome}</b><br>
                {fisc_funcao}
            </div>
        </div>
        """, unsafe_allow_html=True)

        st.markdown("### 8. Download do Documento Pronto")
        col_dl1, col_dl2 = st.columns(2)

        dados_geracao = {
            "numero_req": num_req, "nup": nup_req, "assunto": assunto_req,
            "data_extenso": data_extenso_req, "pregao": pregao_req, "tipo_empenho": tipo_emp_req,
            "razao_social": razao_social_req, "cnpj": cnpj_forn_req, "local_entrega": local_entrega_req,
            "justificativa_a": just_a, "justificativa_b": just_b, "pi": pi_req, "ptres": ptres_req,
            "nd": nd_req, "nc": nc_num_req, "fiscal_nome_posto": fisc_nome, "fiscal_funcao": fisc_funcao
        }

        # Geração em LibreOffice (.odt)
        import zipfile
        buffer_odt = io.BytesIO()
        content_odt_xml = f'''<?xml version="1.0" encoding="UTF-8"?>
    <office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" 
                         xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" 
                         xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" 
                         xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" 
                         office:version="1.2">
      <office:body>
    <office:text>
      <text:h text:outline-level="1">MINISTÉRIO DA DEFESA - EXÉRCITO BRASILEIRO</text:h>
      <text:h text:outline-level="2">BASE DE ADMINISTRAÇÃO E APOIO DA 5ª REGIÃO MILITAR</text:h>
      <text:p>(Companhia do QG da 5ª RM/DI) - BASE MAJOR AGOSTINHO JOSÉ RODRIGUES</text:p>
      <text:p/>
      <text:p>Requisição Nº {num_req} | EB: {nup_req}</text:p>
      <text:p>Assunto: {assunto_req}</text:p>
      <text:p>Curitiba, PR, {data_extenso_req}.</text:p>
      <text:p/>
      <text:p>1. Solicito providências para aprovar a despesa abaixo especificada.</text:p>
      <text:p>2. DEMONSTRATIVO DE NECESSIDADES: {pregao_req}</text:p>
      <text:p>FORNECEDOR: {razao_social_req} (CNPJ: {cnpj_forn_req})</text:p>
      <text:p>VALOR TOTAL: R$ {total_requisicao:,.2f}</text:p>
      <text:p>3. LOCAL DE ENTREGA: {local_entrega_req}</text:p>
      <text:p>4. JUSTIFICATIVA: {just_a}</text:p>
      <text:p>5. CLASSIFICAÇÃO: PI {pi_req} | PTRES {ptres_req} | ND {nd_req} | NC {nc_num_req}</text:p>
      <text:p/>
      <text:p>Aprovado por: {fisc_nome} - {fisc_funcao}</text:p>
    </office:text>
      </office:body>
    </office:document-content>'''
        manifest_odt_xml = '''<?xml version="1.0" encoding="UTF-8"?>
    <manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0">
      <manifest:file-entry manifest:full-path="/" manifest:version="1.2" manifest:media-type="application/vnd.oasis.opendocument.text"/>
      <manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>
    </manifest:manifest>'''
        with zipfile.ZipFile(buffer_odt, 'w', zipfile.ZIP_DEFLATED) as zf:
            zf.writestr('mimetype', 'application/vnd.oasis.opendocument.text', compress_type=zipfile.ZIP_STORED)
            zf.writestr('content.xml', content_odt_xml.encode('utf-8'))
            zf.writestr('META-INF/manifest.xml', manifest_odt_xml.encode('utf-8'))
        buffer_odt.seek(0)

        col_dl1.download_button(
            "📥 Baixar em LibreOffice Writer (.odt)",
            data=buffer_odt,
            file_name=f"Requisicao_{num_req.replace('/', '_').replace(' ', '_')}.odt",
            mime="application/vnd.oasis.opendocument.text",
            use_container_width=True
        )

        # Geração em Word (.docx)
        if docx is not None:
            doc_final = None
            if modelo_id_selecionado:
                c.execute("SELECT arquivo_docx FROM processos_modelos WHERE id = ?", (modelo_id_selecionado,))
                res_bin = c.fetchone()
                if res_bin and res_bin[0]:
                    try:
                        doc_final = docx.Document(io.BytesIO(res_bin[0]))
                    except Exception:
                        doc_final = None

            if doc_final is None:
                doc_final = docx.Document()
                p_c = doc_final.add_paragraph()
                p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r_c = p_c.add_run("MINISTÉRIO DA DEFESA\nEXÉRCITO BRASILEIRO\nBASE DE ADMINISTRAÇÃO E APOIO DA 5ª REGIÃO MILITAR\n(Companhia do QG da 5ª RM/DI)\nBASE MAJOR AGOSTINHO JOSÉ RODRIGUES\n")
                r_c.bold = True
                doc_final.add_paragraph(f"Requisição Nº {num_req}\nEB: {nup_req}\nAssunto: {assunto_req}\nCuritiba, PR, {data_extenso_req}.\n")
                doc_final.add_paragraph(f"DEMONSTRATIVO: {pregao_req}\nFornecedor: {razao_social_req} ({cnpj_forn_req})\nValor: R$ {total_requisicao:,.2f}")
                doc_final.add_paragraph(f"Local de Entrega: {local_entrega_req}\nJustificativa: {just_a}\nClassificação: PI {pi_req} | ND {nd_req} | NC {nc_num_req}")
                p_a = doc_final.add_paragraph(f"\n\n____________________________________________\n{fisc_nome}\n{fisc_funcao}")
                p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER

            buf_docx = io.BytesIO()
            doc_final.save(buf_docx)
            buf_docx.seek(0)

            col_dl2.download_button(
                "📥 Baixar em Word (.docx)",
                data=buf_docx,
                file_name=f"Requisicao_{num_req.replace('/', '_').replace(' ', '_')}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

    with aba_proc2:
        st.markdown("### 📁 Modelos Oficiais Vigentes (LibreOffice e Word)")
        st.caption("Faça o upload dos modelos em formato .odt ou .docx para mantê-los centralizados no banco de dados da OM.")

        c = conn.cursor()
        c.execute("SELECT nome_tipo FROM tipos_documentos WHERE om_id = ? ORDER BY nome_tipo", (user['om_id'],))
        tipos_banco = [r[0] for r in c.fetchall()]

        if not tipos_banco:
            tipos_padrao = [
                "Requisição de Despesa (Pregão Próprio)",
                "Termo de Referência (TR)",
                "Estudo Técnico Preliminar (ETP)",
                "Documento de Formalização da Demanda (DFD)",
                "Pesquisa de Preços / Mapa Comparativo",
                "Justificativa de Adesão (Carona)",
                "Despacho do Ordenador de Despesas"
            ]
            for tp in tipos_padrao:
                c.execute("INSERT OR IGNORE INTO tipos_documentos (om_id, nome_tipo) VALUES (?, ?)", (user['om_id'], tp))
            conn.commit()
            tipos_banco = tipos_padrao

        col_up_m1, col_up_m2 = st.columns(2)
        arq_mod = col_up_m1.file_uploader("Carregar Arquivo de Modelo (.odt ou .docx):", type=["odt", "docx"], key="up_model_file")

        opcoes_tipo_doc = tipos_banco + ["➕ Adicionar Novo Tipo de Documento..."]
        tipo_mod_sel = col_up_m2.selectbox("Tipo de Documento:", opcoes_tipo_doc, key="sel_tipo_mod_box")

        tipo_documento_final = tipo_mod_sel
        if tipo_mod_sel == "➕ Adicionar Novo Tipo de Documento...":
            novo_tipo_digitado = st.text_input("Digite o Nome do Novo Tipo de Documento:", key="input_novo_tipo_doc")
            if novo_tipo_digitado.strip():
                tipo_documento_final = novo_tipo_digitado.strip()

        nome_modelo_custom = st.text_input(
            "Nome de Identificação do Modelo:", 
            value=f"Modelo Padrão - {tipo_documento_final}" if not tipo_documento_final.startswith("➕") else "Novo Modelo",
            key="nome_mod_custom_input"
        )

        if st.button("💾 Salvar Modelo no Banco de Dados", type="primary"):
            if arq_mod is not None:
                if not tipo_documento_final or tipo_documento_final.startswith("➕"):
                    st.error("Por favor, informe um tipo de documento válido.")
                else:
                    bytes_mod = arq_mod.read()
                    extensao = arq_mod.name.split('.')[-1].lower()

                    c.execute("INSERT OR IGNORE INTO tipos_documentos (om_id, nome_tipo) VALUES (?, ?)", (user['om_id'], tipo_documento_final))

                    try:
                        c.execute("ALTER TABLE processos_modelos ADD COLUMN formato TEXT DEFAULT 'odt'")
                    except sqlite3.OperationalError:
                        pass

                    c.execute('''
                    INSERT INTO processos_modelos (om_id, tipo_processo, nome_modelo, arquivo_docx, formato, data_upload)
                    VALUES (?, ?, ?, ?, ?, ?)
                    ''', (user['om_id'], tipo_documento_final, nome_modelo_custom, bytes_mod, extensao, str(datetime.now().date())))
                    conn.commit()

                    st.success(f"✅ Modelo '{nome_modelo_custom}' ({extensao.upper()}) cadastrado com sucesso!")
                    st.rerun()
            else:
                st.warning("Selecione um arquivo .odt ou .docx primeiro.")

        st.markdown("---")
        st.markdown("#### 📑 Modelos Já Cadastrados na OM")

        try:
            c.execute("ALTER TABLE processos_modelos ADD COLUMN formato TEXT DEFAULT 'docx'")
            conn.commit()
        except sqlite3.OperationalError:
            pass

        df_modelos = pd.read_sql_query(f"SELECT id, tipo_processo, nome_modelo, COALESCE(formato, 'docx') as formato, data_upload FROM processos_modelos WHERE om_id = {user['om_id']}", conn)
        if not df_modelos.empty:
            df_modelos['Data Upload'] = df_modelos['data_upload'].apply(formatar_data_br)
            df_modelos['Formato'] = df_modelos['formato'].apply(lambda f: f"📄 LibreOffice ({f.upper()})" if f == 'odt' else f"📄 Word ({f.upper()})")
            st.dataframe(df_modelos[['tipo_processo', 'nome_modelo', 'Formato', 'Data Upload']], use_container_width=True, hide_index=True)
        else:
            st.info("Nenhum modelo cadastrado até o momento.")

    conn.close()
